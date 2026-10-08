"""Build the per-pass Excel report + Allure HTML, and (separately) the one
Xray push, for a run — however many retry attempts each pass actually took.

conftest.py's pytest_sessionfinish only ever dumps its own process's results
to reports/_raw/<label>_attempt<NN>.json (see its module docstring for why).
This script turns those raw dumps, plus the accumulated allure-results/<label>/
for that same label, into that ONE pass's report — retries never add extra
reports, they just update this same one (later attempt's result for a test
overwrites its earlier attempt's result).

There are two reports total for a normal run_suite.sh run — one per pass
(summary_cases, weekly_report_cases) — not one per retry attempt, and not one
combined report across both passes.

Usage:
    python scripts/finalize_report.py --label summary_cases        # that pass's Excel+Allure
    python scripts/finalize_report.py --label weekly_report_cases  # that pass's Excel+Allure
    python scripts/finalize_report.py --xray-only --build 42       # one Xray push, both passes
    python scripts/finalize_report.py                              # standalone/ad hoc: everything
                                                                     # under reports/_raw merged into
                                                                     # one report + one Xray push

Merge rule: raw dumps for the given label are processed in sorted filename
order (attempt number is zero-padded, so this is also attempt order), and a
later attempt's entry for a given test simply overwrites an earlier one — so
the merged result is always each test's LAST attempt. A test whose first
attempt already passed and was never retried keeps that result untouched.
"""
import argparse
import json
import sys
from pathlib import Path

RAW_DIR = Path("reports/_raw")
ALLURE_RESULTS = Path("allure-results")
OUTCOMES_DIR = Path("reports/outcomes")


def load_merged_results(label: str | None = None) -> dict:
    """{(classname, function): result_dict}, last attempt wins per test.

    label=None merges every raw dump regardless of label (used for the
    standalone/ad hoc case, and for the combined Xray push)."""
    pattern = f"{label}_attempt*.json" if label else "*.json"
    merged = {}
    for dump_path in sorted(RAW_DIR.glob(pattern)):
        entries = json.loads(dump_path.read_text())
        for entry in entries:
            merged[(entry["classname"], entry["function"])] = entry
    return merged


def build_excel(results: list[dict], out_path: Path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Results"

    ws.append(["Component", "Type", "Description", "Steps", "Status"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    status_colors = {"PASSED": "92D050", "FAILED": "FF4C4C", "ERROR": "FFA500"}

    for result in results:
        ws.append([
            result["component"],
            result["type"],
            result["description"],
            result["steps"],
            result["status"],
        ])
        last_row = ws.max_row
        color = status_colors.get(result["status"], "FFFFFF")
        ws.cell(last_row, 5).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws.cell(last_row, 4).alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 12

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


def build_allure(alluredir: Path, out_dir: Path):
    import subprocess
    subprocess.run(
        ["allure", "generate", str(alluredir), "-o", str(out_dir), "--clean", "--single-file"],
        check=False,
    )


def push_to_xray(merged: dict, build: str):
    import os
    if not (os.environ.get("XRAY_CLIENT_ID") and os.environ.get("XRAY_CLIENT_SECRET")):
        print("Xray: XRAY_CLIENT_ID/XRAY_CLIENT_SECRET not set — skipping Xray reporting.")
        return

    sys.path.insert(0, str(Path(__file__).parent))
    from xray_common import push_execution_results
    from xray_report import build_entries

    # Xray's status vocabulary isn't pytest's: a skipped test must be reported
    # as "TODO", not the raw "SKIPPED" outcome string (mirrors what
    # xray_common.load_junit_results() has always done for the junit.xml path).
    xray_status = {"SKIPPED": "TODO"}
    junit_results = {
        key: (xray_status.get(e["status"], e["status"]), e["error_detail"])
        for key, e in merged.items()
    }

    print("Xray: matching test results against onboarded keys ...")
    keyed, skipped = build_entries(junit_results=junit_results)
    print(f"Xray: {len(keyed)} keyed result(s) matched, {len(skipped)} unkeyed result(s) skipped.")
    if keyed:
        try:
            execution = push_execution_results(keyed, build)
            print(f"Xray: reported {len(keyed)} result(s) -> {execution['key']}")
        except Exception as exc:
            import traceback
            print(f"Xray: reporting failed, non-fatal ({exc})")
            traceback.print_exc()
    if skipped:
        print(f"Xray: {len(skipped)} test(s) ran without a key, skipped reporting — "
              f"run scripts/xray_onboard.py to onboard them: {', '.join(skipped)}")


def build_report_for_label(label: str | None):
    """label=None covers the standalone/ad hoc case: one report over
    everything under reports/_raw, named by today's date alone."""
    merged = load_merged_results(label)
    if not merged:
        print(f"Nothing to report for label={label!r} — no matching raw dumps under {RAW_DIR}/.")
        return

    results = list(merged.values())
    name = label or "run"
    from datetime import date
    out_dir = OUTCOMES_DIR / f"{name}_{date.today().isoformat()}"
    out_dir.mkdir(parents=True, exist_ok=True)

    alluredir = (ALLURE_RESULTS / label) if label else ALLURE_RESULTS

    # Allure first: `allure generate --clean` clears out_dir before writing
    # its own output, which would otherwise delete the Excel file if it ran
    # second.
    build_allure(alluredir, out_dir)
    print(f"Allure report -> {out_dir}")

    build_excel(results, out_dir / "test_results.xlsx")
    print(f"Excel report -> {out_dir / 'test_results.xlsx'}")


def main():
    import os

    parser = argparse.ArgumentParser()
    parser.add_argument("--label", default=None,
                         help="Build the Excel+Allure report for just this pass (e.g. summary_cases). "
                              "Omit for the standalone/ad hoc case: one report over everything.")
    parser.add_argument("--xray-only", action="store_true",
                         help="Skip Excel/Allure, push every raw dump (any label) to Xray once.")
    parser.add_argument("--build", default=os.environ.get("BUILD_NUMBER", "Local Execution"))
    args = parser.parse_args()

    if not RAW_DIR.exists() or not list(RAW_DIR.glob("*.json")):
        print(f"Nothing to finalize — no raw result dumps found under {RAW_DIR}/. "
              f"Run pytest first (directly, or via scripts/run_suite.sh).")
        return

    if args.xray_only:
        merged = load_merged_results(label=None)
        print(f"Merged {len(merged)} test(s) across all labels for Xray.")
        push_to_xray(merged, args.build)
        return

    build_report_for_label(args.label)
    if args.label is None:
        # Standalone/ad hoc single-shot use also pushes to Xray, same as
        # before this script had --label — keeps `pytest ...` then
        # `python scripts/finalize_report.py` working as a one-step report.
        push_to_xray(load_merged_results(label=None), args.build)


if __name__ == "__main__":
    main()
